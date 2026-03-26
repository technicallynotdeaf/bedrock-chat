"""Document chunking for large text-heavy documents.

When a document attachment exceeds the model's context window capacity,
this module extracts text from the document, splits it into chunks, and
returns the chunks as TextContentModel objects so the model can process
the document content without hitting context limits.

Supported formats: PDF, DOCX, TXT, MD, CSV, HTML, XLSX.
"""

import io
import logging
from typing import Sequence

from app.repositories.models.conversation import (
    AttachmentContentModel,
    ContentModel,
    TextContentModel,
)

logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)

# Approximate token-to-character ratio for English text.
# Claude tokenizes ~4 characters per token on average.
CHARS_PER_TOKEN = 4

# Maximum tokens to allocate for document content in a single message.
# This leaves room for system prompt, conversation history, and model response.
# Bedrock Claude models have 200K context; we reserve ~80K tokens for the document.
MAX_DOCUMENT_TOKENS = 80_000
MAX_DOCUMENT_CHARS = MAX_DOCUMENT_TOKENS * CHARS_PER_TOKEN  # 320,000 chars

# Size threshold in bytes: documents larger than this are candidates for chunking.
# ~500KB of raw bytes (before base64) is roughly where we start risking context overflow.
LARGE_DOCUMENT_THRESHOLD_BYTES = 500_000

# Chunk size in characters for splitting extracted text.
CHUNK_SIZE_CHARS = 50_000  # ~12,500 tokens per chunk


def _extract_text_from_pdf(data: bytes) -> str:
    """Extract text from PDF bytes using pypdf."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"--- Page {i + 1} ---\n{text}")
        return "\n\n".join(pages)
    except Exception as e:
        logger.warning(f"Failed to extract text from PDF: {e}")
        return ""


def _extract_text_from_docx(data: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception as e:
        logger.warning(f"Failed to extract text from DOCX: {e}")
        return ""


def _extract_text_from_xlsx(data: bytes) -> str:
    """Extract text from XLSX bytes using openpyxl."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cell_values = [str(c) if c is not None else "" for c in row]
                if any(v.strip() for v in cell_values):
                    rows.append("\t".join(cell_values))
            if rows:
                sheets.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
        wb.close()
        return "\n\n".join(sheets)
    except Exception as e:
        logger.warning(f"Failed to extract text from XLSX: {e}")
        return ""


def _extract_text_from_plain(data: bytes) -> str:
    """Extract text from plain text formats (txt, md, csv, html)."""
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return ""


def extract_text_from_document(data: bytes, file_name: str) -> str:
    """Extract text content from a document based on its file extension.

    Returns extracted text, or empty string if extraction fails.
    """
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if ext == "pdf":
        return _extract_text_from_pdf(data)
    elif ext in ("docx", "doc"):
        return _extract_text_from_docx(data)
    elif ext in ("xlsx", "xls"):
        return _extract_text_from_xlsx(data)
    elif ext in ("txt", "md", "csv", "html", "css", "json", "py", "js", "ts",
                 "java", "sql", "xml", "yaml", "yml", "ini", "cfg", "log",
                 "sh", "bash", "rb", "go", "rs", "c", "cpp", "h", "hpp"):
        return _extract_text_from_plain(data)
    else:
        # Unsupported format — try plain text as fallback
        text = _extract_text_from_plain(data)
        return text if text and text.isprintable() else ""


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE_CHARS) -> list[str]:
    """Split text into chunks, preferring to break at paragraph boundaries."""
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size

        if end >= len(text):
            chunks.append(text[start:])
            break

        # Try to break at a paragraph boundary (double newline)
        break_pos = text.rfind("\n\n", start, end)
        if break_pos > start + chunk_size // 2:
            end = break_pos + 2  # Include the double newline
        else:
            # Fall back to single newline
            break_pos = text.rfind("\n", start, end)
            if break_pos > start + chunk_size // 2:
                end = break_pos + 1

        chunks.append(text[start:end])
        start = end

    return chunks


def should_chunk_attachment(attachment: AttachmentContentModel) -> bool:
    """Determine if an attachment should be chunked based on size."""
    return len(attachment.body) > LARGE_DOCUMENT_THRESHOLD_BYTES


def chunk_attachment(
    attachment: AttachmentContentModel,
    max_total_chars: int = MAX_DOCUMENT_CHARS,
) -> list[ContentModel]:
    """Convert a large attachment into chunked text content blocks.

    If text extraction succeeds, returns TextContentModel blocks containing
    the document text (truncated to max_total_chars). The original attachment
    is replaced.

    If text extraction fails, returns the original attachment unchanged so
    Bedrock can attempt to process it natively.
    """
    file_name = attachment.file_name
    logger.info(
        f"Attempting to chunk large document: {file_name} "
        f"({len(attachment.body):,} bytes)"
    )

    extracted_text = extract_text_from_document(attachment.body, file_name)

    if not extracted_text.strip():
        logger.warning(
            f"Could not extract text from {file_name}. "
            f"Sending original attachment to Bedrock."
        )
        return [attachment]

    logger.info(
        f"Extracted {len(extracted_text):,} characters from {file_name} "
        f"(~{len(extracted_text) // CHARS_PER_TOKEN:,} tokens)"
    )

    # Truncate if needed
    if len(extracted_text) > max_total_chars:
        logger.info(
            f"Truncating document from {len(extracted_text):,} to "
            f"{max_total_chars:,} characters to fit context window"
        )
        # Find a clean break point near the limit
        break_pos = extracted_text.rfind("\n\n", 0, max_total_chars)
        if break_pos < max_total_chars // 2:
            break_pos = extracted_text.rfind("\n", 0, max_total_chars)
        if break_pos < max_total_chars // 2:
            break_pos = max_total_chars
        extracted_text = extracted_text[:break_pos]

    chunks = chunk_text(extracted_text)
    total_chunks = len(chunks)

    result: list[ContentModel] = []

    # Header block
    header = (
        f"[Document: {file_name} — Text extracted and split into "
        f"{total_chunks} chunk(s) because the original document was too large "
        f"for the context window]"
    )

    for i, chunk in enumerate(chunks):
        chunk_label = f"\n\n--- {file_name} (chunk {i + 1}/{total_chunks}) ---\n\n"
        body = (header + chunk_label + chunk) if i == 0 else (chunk_label + chunk)
        result.append(TextContentModel(content_type="text", body=body))

    logger.info(f"Split {file_name} into {total_chunks} text chunk(s)")
    return result


def process_attachments_for_context_window(
    content: list[ContentModel],
) -> list[ContentModel]:
    """Process message content, chunking any oversized attachments.

    Iterates through content blocks. Any AttachmentContentModel that exceeds
    the size threshold is extracted to text and chunked. Other content blocks
    pass through unchanged.

    Returns a new content list with large attachments replaced by text chunks.
    """
    result: list[ContentModel] = []
    chunked_any = False

    for item in content:
        if isinstance(item, AttachmentContentModel) and should_chunk_attachment(item):
            chunked_content = chunk_attachment(item)
            result.extend(chunked_content)
            if not (len(chunked_content) == 1 and chunked_content[0] is item):
                chunked_any = True
        else:
            result.append(item)

    if chunked_any:
        logger.info(
            f"Processed message content: {len(content)} blocks -> {len(result)} blocks"
        )

    return result
